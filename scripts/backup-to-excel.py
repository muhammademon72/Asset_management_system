"""
Convert assets-backup JSON to a readable Excel workbook.
Each collection becomes a separate sheet.
"""
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BACKUP_PATH = "/home/z/my-project/upload/assets-backup-20261026-075311.json"
OUT_PATH = "/home/z/my-project/download/assets-backup-2026-10-26.xlsx"

# Load backup
with open(BACKUP_PATH, "r", encoding="utf-8") as f:
    bk = json.load(f)

collections = bk.get("collections", {})

wb = Workbook()
wb.remove(wb.active)

# Styles
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="B0B0B0")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ZEBRA_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUB_FONT = Font(italic=True, size=10, color="666666")

# --- Sheet 0: Summary ---
ws = wb.create_sheet("Summary", 0)
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 22
ws["A1"] = "Asset Management System — Backup Summary"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A1:C1")
ws["A2"] = "Exported: " + bk.get("exportedAt", "?")
ws["A2"].font = SUB_FONT
ws.merge_cells("A2:C2")
ws["A3"] = "Version: " + str(bk.get("version", "?"))
ws["A3"].font = SUB_FONT
ws.merge_cells("A3:C3")

ws["A5"] = "Collection"
ws["B5"] = "Records"
ws["C5"] = "Notes"
for c in ("A5", "B5", "C5"):
    ws[c].font = HEADER_FONT
    ws[c].fill = HEADER_FILL
    ws[c].alignment = HEADER_ALIGN
    ws[c].border = CELL_BORDER

notes_map = {
    "branches": "All branch locations",
    "employees": "Employee master records",
    "departments": "Department list",
    "designations": "Job titles",
    "products": "Product catalog",
    "categories": "Product categories",
    "brands": "Brand list",
    "uoms": "Units of measure",
    "conditions": "Item conditions (New/Used/Refurbished)",
    "warranties": "Warranty periods",
    "suppliers": "Supplier master",
    "purchases": "Purchase invoices — multi-product rows",
    "deliveries": "Delivery challans — multi-product rows",
    "adjustments": "Inventory adjustments",
    "damages": "Damaged items",
    "transfers": "Stock transfers",
    "cashIn": "Cash receipts",
    "cashOut": "Cash payments",
    "users": "Login accounts",
    "roles": "Role & privilege definitions",
    "company": "Company profile (single record)",
}

r = 6
total = 0
for name, items in collections.items():
    if not isinstance(items, list):
        continue
    cnt = len(items)
    total += cnt
    ws.cell(row=r, column=1, value=name).border = CELL_BORDER
    ws.cell(row=r, column=2, value=cnt).border = CELL_BORDER
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3, value=notes_map.get(name, "")).border = CELL_BORDER
    if (r - 6) % 2 == 1:
        for col in (1, 2, 3):
            ws.cell(row=r, column=col).fill = ZEBRA_FILL
    r += 1

# Total row
ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor="D9E1F2")
ws.cell(row=r, column=2, value=total).font = Font(bold=True)
ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="D9E1F2")
ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
ws.cell(row=r, column=3, value="Total records across all collections")
ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor="D9E1F2")

# --- One sheet per collection ---
SKIP_KEYS = {"attachment"}  # skip attachment base64 blobs (too large)

def write_collection_sheet(name, items):
    if not isinstance(items, list) or not items:
        # create an empty sheet so user sees it exists
        s = wb.create_sheet(name[:31])
        s["A1"] = "(no records)"
        s["A1"].font = SUB_FONT
        return
    s = wb.create_sheet(name[:31])
    # collect all keys (preserve insertion order across rows)
    keys = []
    seen = set()
    for item in items:
        if not isinstance(item, dict):
            continue
        for k in item.keys():
            if k in SKIP_KEYS:
                continue
            if k not in seen:
                seen.add(k)
                keys.append(k)
    # header row
    for ci, key in enumerate(keys, start=1):
        c = s.cell(row=1, column=ci, value=key)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN
        c.border = CELL_BORDER
    # data rows
    ri = 2
    for item in items:
        if not isinstance(item, dict):
            s.cell(row=ri, column=1, value=str(item))
            ri += 1
            continue
        for ci, key in enumerate(keys, start=1):
            v = item.get(key, "")
            # flatten nested objects/dicts to JSON
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False)
            elif isinstance(v, str) and len(v) > 1000:
                v = v[:1000] + "...[truncated]"
            c = s.cell(row=ri, column=ci, value=v)
            c.border = CELL_BORDER
            c.alignment = Alignment(vertical="top", wrap_text=False)
            if (ri - 2) % 2 == 1:
                c.fill = ZEBRA_FILL
        ri += 1
    # auto column widths
    for ci, key in enumerate(keys, start=1):
        max_len = len(str(key))
        # sample first 50 rows
        for ri2 in range(2, min(ri, 52)):
            cv = s.cell(row=ri2, column=ci).value
            if cv is not None:
                ln = len(str(cv))
                if ln > max_len:
                    max_len = ln
        # cap width
        col_w = min(max(max_len + 2, 10), 50)
        s.column_dimensions[get_column_letter(ci)].width = col_w
    # freeze header
    s.freeze_panes = "A2"
    # autofilter
    s.auto_filter.ref = "A1:" + get_column_letter(len(keys)) + str(max(ri - 1, 2))

# write in same order as Summary
for name in collections.keys():
    write_collection_sheet(name, collections[name])

wb.properties.creator = "Z.ai"
wb.properties.title = "Asset Management System Backup"
wb.save(OUT_PATH)
print(f"OK: {OUT_PATH}")
print(f"Sheets: {len(wb.sheetnames)}")
print("Sheet names:", wb.sheetnames)
print(f"File size: {os.path.getsize(OUT_PATH)/1024:.1f} KB")
